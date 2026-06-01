"""端到端测试：扫描 → 解析 → PDF判定 完整流程。"""
import os
import json
import tempfile
import pytest
from app.engine.scanner import scan_directory
from app.engine.file_parser import parse_file
from app.engine.pdf_judge import judge_pdf


@pytest.fixture
def test_workspace():
    tmp = tempfile.mkdtemp()
    # 创建 txt
    with open(os.path.join(tmp, "readme.txt"), "w", encoding="utf-8") as f:
        f.write("Hello world\n测试中文")
    # 创建 csv
    with open(os.path.join(tmp, "data.csv"), "w", encoding="utf-8") as f:
        f.write("name,age\n张三,30")
    # 创建 json
    with open(os.path.join(tmp, "config.json"), "w", encoding="utf-8") as f:
        json.dump({"app": "DocScanX"}, f)
    # 创建 docx
    from docx import Document
    doc = Document()
    doc.add_paragraph("员工合同文本")
    doc.save(os.path.join(tmp, "contract.docx"))
    # 创建 xlsx
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["姓名", "手机号"])
    ws.append(["张三", "13800001111"])
    wb.save(os.path.join(tmp, "roster.xlsx"))

    yield tmp
    import shutil
    shutil.rmtree(tmp)


class TestPipeline:
    def test_scan_then_parse_all(self, test_workspace):
        groups = {"text": True, "office_new": True, "pdf": True, "structured": True, "archive": False}
        files = scan_directory(test_workspace, ext_groups=groups)

        assert len(files) >= 5  # txt + csv + docx + xlsx + json

        errors = []
        for f in files:
            result = parse_file(f.path)
            if result.error:
                errors.append((f.ext, result.error))

        assert errors == [], f"Parse errors: {errors}"

        # 至少有一个文件包含中文
        all_text = " ".join(parse_file(f.path).text for f in files)
        assert "张三" in all_text

        print(f"Pipeline OK: {len(files)} files parsed")

    def test_pdf_judge_on_non_pdf_wont_crash(self):
        """pdf_judge 在非 PDF 上应返回错误，不抛异常。"""
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w") as f:
            f.write("not a pdf")
            tmp_path = f.name
        try:
            result = judge_pdf(tmp_path)
            assert isinstance(result.error, str)
            assert result.error != ""
        finally:
            os.unlink(tmp_path)
